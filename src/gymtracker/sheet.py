"""Workouts.xlsx, rebuilt from the database every time (DESIGN.md §4.5).

Two tabs: Workouts (one row per workout) and Monthly (totals per month). Values are real
Excel dates, times and numbers, so they sort, filter and chart in Excel.

Kept short on purpose (user, 2026-10-18): no wall-clock duration, incline, ascent, pauses,
calories or source (new workouts all come from the tracker). All of those are still in the
database and archive, so a column can come back by adding it here.
"""

from __future__ import annotations

import os
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from gymtracker import units
from gymtracker.store import Workout

_DAY_S = 86400


@dataclass(frozen=True)
class _Units:
    dist: str
    speed: str
    pace: str
    m_to_dist: float


IMPERIAL = _Units("mi", "mph", "min/mi", 1 / units.M_PER_MILE)
METRIC = _Units("km", "km/h", "min/km", 1 / 1000)


def _units(name: str) -> _Units:
    return METRIC if name == "metric" else IMPERIAL


def _days(seconds: float | None) -> float | None:
    """Excel stores durations as fractions of a day."""
    return None if seconds is None else seconds / _DAY_S


def _speed(mps: float | None, u: _Units) -> float | None:
    return None if mps is None else mps * 3600 * u.m_to_dist


def _pace_days(distance_m: float, moving_s: float, u: _Units) -> float | None:
    dist = distance_m * u.m_to_dist
    return _days(moving_s / dist) if dist > 0 and moving_s > 0 else None


def build_workbook(workouts: Iterable[Workout], unit_system: str = "imperial") -> Workbook:
    u = _units(unit_system)
    workouts = sorted(workouts, key=lambda w: w.summary.start_utc)
    wb = Workbook()
    ws = wb.active
    ws.title = "Workouts"
    columns = [
        ("Date", 16, "ddd m/d/yyyy"),
        ("Start", 10, "h:mm AM/PM"),
        ("Moving", 10, "[h]:mm:ss"),
        (f"Distance ({u.dist})", 13, "0.00"),
        (f"Pace ({u.pace})", 13, "mm:ss"),
        (f"Min speed ({u.speed})", 16, "0.0"),
        (f"Avg speed ({u.speed})", 16, "0.0"),
        (f"Max speed ({u.speed})", 16, "0.0"),
        ("Min HR", 9, "0"),
        ("Avg HR", 9, "0"),
        ("Max HR", 9, "0"),
        ("Notes", 40, "@"),
    ]
    _header(ws, columns)
    for w in workouts:
        s = w.summary
        start = w.start_local
        ws.append([
            start.date(), start.time().replace(tzinfo=None), _days(s.moving_s),
            s.distance_m * u.m_to_dist, _pace_days(s.distance_m, s.moving_s, u),
            _speed(s.min_speed_mps, u), _speed(s.avg_speed_mps, u), _speed(s.max_speed_mps, u),
            s.min_hr, s.avg_hr, s.max_hr, w.notes,
        ])
    _formats(ws, columns)

    ms = wb.create_sheet("Monthly")
    mcols = [
        ("Month", 12, "mmm yyyy"),
        ("Workouts", 10, "0"),
        ("Moving", 10, "[h]:mm"),
        (f"Distance ({u.dist})", 13, "0.0"),
        (f"Avg pace ({u.pace})", 16, "mm:ss"),
        ("Avg HR", 9, "0"),
    ]
    _header(ms, mcols)
    for month, group in _by_month(workouts):
        moving = sum(w.summary.moving_s for w in group)
        dist = sum(w.summary.distance_m for w in group)
        hr_w = [(w.summary.avg_hr, w.summary.moving_s) for w in group if w.summary.avg_hr]
        hr_time = sum(t for _, t in hr_w)
        avg_hr = round(sum(h * t for h, t in hr_w) / hr_time) if hr_time else None
        ms.append([month, len(group), _days(moving), dist * u.m_to_dist,
                   _pace_days(dist, moving, u), avg_hr])
    _formats(ms, mcols)
    return wb


def write_workbook(workouts: Iterable[Workout], path: str | Path,
                   unit_system: str = "imperial") -> Path:
    """Build and save atomically (a reader never sees a half-written file)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".tmp" + path.suffix)
    build_workbook(workouts, unit_system).save(tmp)
    os.replace(tmp, path)
    return path


def _by_month(workouts: list[Workout]) -> list[tuple[date, list[Workout]]]:
    groups: dict[date, list[Workout]] = {}
    for w in workouts:
        start = w.start_local
        groups.setdefault(date(start.year, start.month, 1), []).append(w)
    return sorted(groups.items())


def _header(ws, columns) -> None:
    ws.append([name for name, _, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _formats(ws, columns) -> None:
    for i, (_, width, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.number_format = fmt
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
