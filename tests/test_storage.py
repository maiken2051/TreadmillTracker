"""Storage: database, archive, library, spreadsheet, .fit import, rebuild."""

import gzip
import json
from datetime import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from gymtracker.archive import read_archive, write_archive
from gymtracker.config import SessionConfig
from gymtracker.library import Library, rebuild_db
from gymtracker.recorder import Recorder
from gymtracker.session import Record, summarize
from gymtracker.sheet import build_workbook
from gymtracker.sim.replay import read_fixture, run_recorder
from gymtracker.store import DuplicateWorkout, Store, Workout

FIXTURES = Path(__file__).parent / "fixtures"
EDT = -240  # minutes


def make_workout(start_t=1_760_000_000.0, n=120, hr=100, source="pi", offset=EDT):
    records = [Record(t, hr + t % 7, 1.2, 1.2 * t, 1.0, t / 10, False) for t in range(n)]
    summary = summarize(start_t, start_t + n, records, distance_m=1.2 * (n - 1),
                        kcal=(n - 1) / 10)
    return Workout(summary=summary, utc_offset_min=offset, source=source,
                   devices={"treadmill": "Treadmill 0001"}), records


def saved_calib_session():
    saved = []
    notes = read_fixture(FIXTURES / "treadmill_calib.jsonl")
    run_recorder(notes, Recorder(SessionConfig(), on_save=saved.append))
    raw = [{"t": round(n.t * 1000), "dev": n.dev, "char": n.char, "data": n.data.hex()}
           for n in notes]
    return saved[0], raw


# --------------------------------------------------------------------------- store


def test_store_round_trip(tmp_path):
    w, records = make_workout()
    with Store(tmp_path / "db.sqlite") as st:
        wid = st.save(w, records)
        got = st.get(wid)
        assert got.summary == w.summary
        assert got.devices == w.devices and got.source == "pi" and got.created_utc
        assert st.samples(wid) == records
        assert [x.id for x in st.list()] == [wid]


def test_store_rejects_duplicate_start(tmp_path):
    w, records = make_workout()
    with Store(tmp_path / "db.sqlite") as st:
        st.save(w, records)
        with pytest.raises(DuplicateWorkout):
            st.save(make_workout()[0], records)


def test_store_delete_cascades_and_notes(tmp_path):
    w, records = make_workout()
    with Store(tmp_path / "db.sqlite") as st:
        wid = st.save(w, records)
        st.set_notes(wid, "felt good")
        assert st.get(wid).notes == "felt good"
        st.delete(wid)
        assert st.get(wid) is None
        assert st.db.execute("SELECT COUNT(*) FROM samples").fetchone()[0] == 0


def test_store_reopens_and_snapshots(tmp_path):
    w, records = make_workout()
    with Store(tmp_path / "db.sqlite") as st:
        st.save(w, records)
        st.snapshot_to(tmp_path / "copy.sqlite")
    with Store(tmp_path / "db.sqlite") as st:
        assert len(st.list()) == 1
    with Store(tmp_path / "copy.sqlite") as st:
        assert len(st.list()) == 1


def test_migration_from_schema_1_backfills_min_speed(tmp_path):
    import sqlite3

    from gymtracker import store as store_mod
    w, records = make_workout(n=200)
    with Store(tmp_path / "db.sqlite") as st:
        wid = st.save(w, records)
    db = sqlite3.connect(tmp_path / "db.sqlite")  # turn it back into a schema 1 database
    db.execute("ALTER TABLE workouts DROP COLUMN min_speed_mps")
    db.execute("PRAGMA user_version = 1")
    db.commit()
    db.close()
    with Store(tmp_path / "db.sqlite") as st:
        assert st.db.execute("PRAGMA user_version").fetchone()[0] == store_mod.SCHEMA_VERSION
        assert st.get(wid).summary.min_speed_mps == pytest.approx(1.2)


def test_store_refuses_newer_schema(tmp_path):
    with Store(tmp_path / "db.sqlite") as st:
        st.db.execute("PRAGMA user_version = 99")
    with pytest.raises(RuntimeError):
        Store(tmp_path / "db.sqlite")


def test_local_start_uses_offset():
    w, _ = make_workout(start_t=1_760_660_575.0)  # 2025-10-17 00:22:55 UTC
    assert w.start_local.strftime("%Y-%m-%d %H:%M") == "2025-10-16 20:22"


# --------------------------------------------------------------------------- archive


def test_archive_round_trip_and_name(tmp_path):
    w, records = make_workout(start_t=1_760_660_575.0)
    raw = [{"t": 1, "dev": "x", "char": "y", "data": "00"}]
    path = write_archive(tmp_path, w, records, raw)
    assert path.relative_to(tmp_path).as_posix() == "2025/2025-10-16_20-22_treadmill.json.gz"
    a = read_archive(path)
    assert a.workout.summary == w.summary and a.records == records and a.raw == raw
    with gzip.open(path, "rt") as f:
        assert json.load(f)["format"] == "gymtracker-workout"


def test_archive_never_overwrites(tmp_path):
    w, records = make_workout()
    first = write_archive(tmp_path, w, records)
    second = write_archive(tmp_path, w, records)
    third = write_archive(tmp_path, w, records)
    assert len({first, second, third}) == 3
    assert second.name.endswith("_treadmill-2.json.gz")
    assert third.name.endswith("_treadmill-3.json.gz")


def test_archive_without_min_speed_recomputes_it(tmp_path):
    w, records = make_workout(n=200)
    path = write_archive(tmp_path, w, records)
    with gzip.open(path, "rt") as f:
        doc = json.load(f)
    del doc["workout"]["summary"]["min_speed_mps"]
    with gzip.open(path, "wt") as f:
        json.dump(doc, f)
    assert read_archive(path).workout.summary.min_speed_mps == pytest.approx(1.2)


def test_archive_rejects_other_files(tmp_path):
    p = tmp_path / "x.json.gz"
    with gzip.open(p, "wt") as f:
        json.dump({"format": "other"}, f)
    with pytest.raises(ValueError):
        read_archive(p)


# --------------------------------------------------------------------------- library


def test_library_add_session_writes_archive_and_row(tmp_path):
    session, raw = saved_calib_session()
    with Library(tmp_path) as lib:
        wid = lib.add_session(session, raw)
        w = lib.store.get(wid)
        assert (w.summary.moving_s, w.summary.distance_m, w.summary.kcal) == (308, 418, 29)
        assert (lib.archive_dir / w.archive).exists()
        assert len(read_archive(lib.archive_dir / w.archive).raw) == len(raw)
        with pytest.raises(DuplicateWorkout):
            lib.add_session(session, raw)


def test_library_delete_moves_archive_aside(tmp_path):
    w, records = make_workout()
    with Library(tmp_path) as lib:
        wid = lib.add(w, records)
        rel = lib.store.get(wid).archive
        lib.delete(wid)
        assert lib.store.get(wid) is None
        assert not (lib.archive_dir / rel).exists()
        assert (lib.archive_dir / "deleted" / rel).exists()
    count, problems = rebuild_db(tmp_path / "archive", tmp_path / "new.db")
    assert count == 0 and problems == []


def test_rebuild_gives_identical_database(tmp_path):
    session, raw = saved_calib_session()
    with Library(tmp_path) as lib:
        lib.add_session(session, raw)
        for i in range(3):
            lib.add(*make_workout(start_t=1_750_000_000.0 + i * 86400, source="import:x"))
        lib.store.set_notes(1, "notes live in the database only")
        before = lib.store.list()
        samples = {w.summary.start_utc: lib.store.samples(w.id) for w in before}
    count, problems = rebuild_db(tmp_path / "archive", tmp_path / "rebuilt.db")
    assert count == 4 and problems == []
    with Store(tmp_path / "rebuilt.db") as st:
        after = st.list()
        assert [w.summary for w in after] == [w.summary for w in before]
        assert [w.archive for w in after] == [w.archive for w in before]
        assert {w.summary.start_utc: st.samples(w.id) for w in after} == samples
    with pytest.raises(FileExistsError):
        rebuild_db(tmp_path / "archive", tmp_path / "rebuilt.db")


def test_rebuild_reparse_uses_raw(tmp_path):
    session, raw = saved_calib_session()
    with Library(tmp_path) as lib:
        lib.add_session(session, raw)
    count, problems = rebuild_db(tmp_path / "archive", tmp_path / "re.db", reparse=True)
    assert count == 1 and problems == []
    with Store(tmp_path / "re.db") as st:
        assert st.list()[0].summary == session.summary()


# --------------------------------------------------------------------------- spreadsheet


def test_workbook_rows_and_units(tmp_path):
    session, raw = saved_calib_session()
    with Library(tmp_path) as lib:
        lib.add_session(session, raw)
        lib.add(*make_workout(start_t=1_760_660_575.0, n=600))
        path = lib.write_sheet()
    wb = load_workbook(path)
    ws = wb["Workouts"]
    header = [c.value for c in ws[1]]
    assert header == ["Date", "Start", "Moving", "Distance (mi)", "Pace (min/mi)",
                      "Min speed (mph)", "Avg speed (mph)", "Max speed (mph)", "Min HR",
                      "Avg HR", "Max HR", "Notes"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    first = dict(zip(header, rows[0], strict=True))
    assert first["Start"] == time(20, 22, 55)
    assert first["Distance (mi)"] == pytest.approx(1.2 * 599 / 1609.344)
    assert first["Min HR"] <= first["Avg HR"] <= first["Max HR"]
    calib = dict(zip(header, rows[1], strict=True))
    assert round(calib["Distance (mi)"], 2) == 0.26
    assert round(calib["Moving"].total_seconds()) == 308  # a real Excel duration
    # calibration walk: 2.0 mph (68 s), 3.0 (116 s), 4.0 (60 s)
    assert (round(calib["Min speed (mph)"], 1), round(calib["Max speed (mph)"], 1)) == (2.0, 4.0)
    monthly = list(wb["Monthly"].iter_rows(min_row=2, values_only=True))
    assert [m[1] for m in monthly] == [1, 1]  # Oct 2025, Oct 2026


def test_workbook_metric(tmp_path):
    wb = build_workbook([make_workout(n=1001)[0]], "metric")
    ws = wb["Workouts"]
    assert ws.cell(1, 4).value == "Distance (km)"
    assert ws.cell(2, 4).value == pytest.approx(1.2)


def test_write_sheet_is_atomic(tmp_path):
    with Library(tmp_path) as lib:
        path = lib.write_sheet()
    assert path.exists() and not list(tmp_path.glob("*.tmp*"))


# --------------------------------------------------------------------------- .fit import


def _write_fit(path, start_ms, seconds, hr=110, per_second=3, timer_s=None):
    from fit_tool.fit_file_builder import FitFileBuilder
    from fit_tool.profile.messages.file_id_message import FileIdMessage
    from fit_tool.profile.messages.record_message import RecordMessage
    from fit_tool.profile.messages.session_message import SessionMessage
    from fit_tool.profile.profile_type import FileType, Manufacturer, Sport, SubSport

    b = FitFileBuilder(auto_define=True)
    fid = FileIdMessage()
    fid.type = FileType.ACTIVITY
    fid.manufacturer = Manufacturer.DEVELOPMENT.value
    fid.time_created = start_ms
    b.add(fid)
    for t in range(seconds):
        for _ in range(per_second):
            r = RecordMessage()
            r.timestamp = start_ms + t * 1000
            r.heart_rate = 0 if t == 0 else hr
            r.distance = 1.0 * t
            r.enhanced_speed = 1.0
            r.calories = t // 10
            b.add(r)
    s = SessionMessage()
    s.timestamp = start_ms + (seconds - 1) * 1000
    s.start_time = start_ms
    s.sport = Sport.RUNNING
    s.sub_sport = SubSport.INDOOR_RUNNING
    s.total_elapsed_time = timer_s or seconds
    s.total_distance = float(seconds - 1)
    s.total_calories = (seconds - 1) // 10
    b.add(s)
    b.build().to_file(str(path))


def test_import_fit_collapses_to_1hz_and_dedupes(tmp_path):
    src = tmp_path / "fits"
    src.mkdir()
    _write_fit(src / "a.fit", 1_760_660_575_000, 300, hr=110, timer_s=280)
    _write_fit(src / "b.FIT", 1_760_760_575_000, 120, hr=95)
    (src / "notes.txt").write_text("ignored")
    with Library(tmp_path / "lib") as lib:
        result = lib.import_fit(src, "test")
        assert len(result.added) == 2 and not result.failed
        a = lib.store.list()[0]
        assert a.source == "import:test"
        assert a.summary.moving_s == 280  # the file's own session time
        assert a.summary.distance_m == 299
        assert a.summary.avg_hr == 110  # the HR 0 second is a gap, not a zero
        samples = lib.store.samples(a.id)
        assert len(samples) == 300 and samples[0].hr is None
        again = lib.import_fit(src, "test")
        assert not again.added and len(again.duplicates) == 2


def test_import_fit_reports_bad_files(tmp_path):
    src = tmp_path / "fits"
    src.mkdir()
    (src / "broken.fit").write_bytes(b"not a fit file")
    with Library(tmp_path / "lib") as lib:
        result = lib.import_fit(src, "test")
    assert not result.added and len(result.failed) == 1
